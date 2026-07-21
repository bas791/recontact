#!/usr/bin/env python3
"""Extract the recontact customer list from a ServiceM8 completed-jobs export.

Reads the "Jan-Dec 2024" sheet, groups jobs by customer email, and assigns
each customer to a salesperson from the row highlight colour:

    green (FFC6EFCE)  -> Bas
    red   (FFFFC7CE)  -> Herman
    no colour         -> Herman

If a customer has jobs in more than one colour, green wins (Bas), then red.

"Last serviced" is the latest "Scheduled for ... on d/m/yyyy" date found in
the Booking column (falling back to payment dates), capped at today — the
booking history runs well past the 2024 completion dates and is the honest
signal for how recently we were on site.

Usage:
    python3 scripts/extract_customers.py <export.xlsx> [out.json]
"""
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl

GREEN = "FFC6EFCE"
RED = "FFFFC7CE"
SHEET = "Jan-Dec 2024"

SCHED_RE = re.compile(r"Scheduled for .{0,300}? on (\d{1,2})/(\d{1,2})/(\d{4})", re.S)
DATE_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
BAD_PHONES = {"NA", "N/A", "-", "NONE", "TBC"}


def parse_dates(matches, today):
    out = []
    for d, m, y in matches:
        try:
            dt = datetime(int(y), int(m), int(d))
        except ValueError:
            continue
        if datetime(2023, 1, 1) <= dt <= today:
            out.append(dt)
    return out


def money(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = re.sub(r"[^\d.]", "", v)
        return float(s) if s else 0.0
    return 0.0


def main(xlsx_path, out_path="data/customers.json"):
    today = datetime.now()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_fmt = openpyxl.load_workbook(xlsx_path)
    ws, ws_fmt = wb[SHEET], wb_fmt[SHEET]

    def row_colour(r):
        for c in ws_fmt[r]:
            f = c.fill
            if f and f.fill_type == "solid" and f.start_color.type == "rgb":
                if f.start_color.rgb == GREEN:
                    return "green"
                if f.start_color.rgb == RED:
                    return "red"
        return "none"

    jobs = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 3).value
        if not name:
            continue
        g = lambda c: str(ws.cell(r, c).value or "")
        sched = parse_dates(SCHED_RE.findall(g(11)), today)
        pay = parse_dates(DATE_RE.findall(g(10)), today)
        any_k = parse_dates(DATE_RE.findall(g(11)), today)
        jd = max(sched or pay or any_k, default=None)
        phone = g(6).strip()
        email = g(7).strip().lower()
        if "@" not in email:  # placeholder like "na" — not a real address
            email = ""
        jobs.append(dict(
            jobNumber=g(2).strip(),
            name=str(name).strip(),
            clientType=g(4).strip(),
            site=g(5).strip(),
            phone="" if phone.upper() in BAD_PHONES else phone,
            email=email,
            desc=g(9).strip(),
            value=money(ws.cell(r, 15).value),
            date=jd.strftime("%Y-%m-%d") if jd else None,
            cat=row_colour(r),
        ))

    # group by email; rows with no real email group by customer name instead
    by_key = defaultdict(list)
    for j in jobs:
        key = j["email"] or "name:" + re.sub(r"\s+", " ", j["name"].lower())
        by_key[key].append(j)

    customers = []
    for key, js in by_key.items():
        email = js[0]["email"]
        js.sort(key=lambda j: j["date"] or "0000", reverse=True)
        names = Counter(j["name"] for j in js)
        types = Counter(j["clientType"] for j in js if j["clientType"])
        cats = {j["cat"] for j in js}
        customers.append(dict(
            key=key,
            email=email,
            name=names.most_common(1)[0][0],
            clientType=types.most_common(1)[0][0] if types else "",
            phone=next((j["phone"] for j in js if j["phone"]), ""),
            owner="bas" if "green" in cats else "herman",
            colour="green" if "green" in cats else ("red" if "red" in cats else "none"),
            totalValue=round(sum(j["value"] for j in js), 2),
            jobCount=len(js),
            lastJobDate=js[0]["date"],
            jobs=[dict(
                jobNumber=j["jobNumber"],
                desc=" ".join(j["desc"].split())[:300],
                value=j["value"],
                date=j["date"],
                site=j["site"].split("\n")[0][:80],
            ) for j in js],
        ))

    customers.sort(key=lambda c: -c["totalValue"])
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(customers, f, ensure_ascii=False)

    bas = [c for c in customers if c["owner"] == "bas"]
    herman = [c for c in customers if c["owner"] == "herman"]
    print(f"{len(customers)} customers -> {out_path}")
    print(f"  Bas:    {len(bas)} (${sum(c['totalValue'] for c in bas):,.0f})")
    print(f"  Herman: {len(herman)} (${sum(c['totalValue'] for c in herman):,.0f})")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(*sys.argv[1:3])
